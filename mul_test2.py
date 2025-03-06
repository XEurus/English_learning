import json
import spacy
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
import multiprocessing

def init_pool(shared_word_list, shared_word_info):
    global word_set
    global word_info_dict
    global nlp
    word_set = set(shared_word_list)
    word_info_dict = shared_word_info
    nlp = spacy.load('en_core_web_sm')

def extract_headwords_and_info_from_file(file_path):
    """
    从文件中逐行提取JSON数据，提取headWord字段和对应的信息。

    参数:
    file_path (str): 文件路径，文件包含多个JSON块。

    返回:
    tuple: (headwords, word_info_dict)
        - headwords: 包含所有提取的headWord的列表。
        - word_info_dict: 以headWord为键，词义和其他信息为值的字典。
    """
    headwords = []
    word_info_dict = {}

    try:
        with open(file_path, 'r', encoding='utf-8') as file:
            for line in file:
                # 去除可能的空行和首尾空白
                line = line.strip()
                if line:  # 只有非空行才处理
                    try:
                        data = json.loads(line)  # 将每一行解析为JSON对象
                        head_word = data['headWord']
                        # 提取词义和其他信息（根据您的 JSON 结构调整）
                        content = data.get('content', {})
                        word_info = content.get('word', {}).get('content', {})
                        # 提取释义
                        translation = word_info.get('trans', [])
                        if translation:
                            trans_text = '; '.join([tran.get('tranCn', '') for tran in translation])
                        else:
                            trans_text = ''
                        word_info_dict[head_word.lower()] = trans_text
                        headwords.append(head_word)
                    except (json.JSONDecodeError, KeyError) as e:
                        print(f"Error processing line: {line}\nError: {e}")
                        continue  # 跳过错误的行

    except FileNotFoundError:
        print(f"File not found: {file_path}")

    return headwords, word_info_dict

def process_paragraph(paragraph_text):
    global word_set, word_info_dict, nlp
    new_runs = []
    found_words = set()

    # 以换行符为分隔，拆分段落为句子
    sentences = paragraph_text.split('\n')

    for sentence_text in sentences:
        sentence_found_words = set()
        sentence_runs = []

        # 使用 spaCy 处理每个句子，以便进行词形还原和标记化
        doc = nlp(sentence_text)

        for token in doc:
            text = token.text_with_ws
            word = token.lemma_.lower()
            if word in word_set:
                found_words.add(word)
                sentence_found_words.add(word)
                run_info = {
                    'text': text,
                    'highlight': True
                }
            else:
                run_info = {
                    'text': text,
                    'highlight': False
                }
            sentence_runs.append(run_info)

        # 将句子的 runs 添加到 new_runs
        new_runs.extend(sentence_runs)

        # 如果该句子中有高亮的单词，在句子后面添加释义
        if sentence_found_words:
            # 组装释义信息
            meanings = []
            for fw in sentence_found_words:
                meaning = word_info_dict.get(fw, '')
                if meaning:
                    meanings.append(f"{fw}: {meaning}")
                else:
                    meanings.append(f"{fw}")
            # 创建新的 run，包含释义信息
            meanings_text = ' [' + '; '.join(meanings) + ']'
            new_runs.append({
                'text': meanings_text,
                'highlight': False
            })

        # 添加换行符，以保持原有的段落结构
        new_runs.append({
            'text': '\n',
            'highlight': False
        })

    # 去除最后多余的换行符
    if new_runs and new_runs[-1]['text'] == '\n':
        new_runs.pop()

    return new_runs, found_words

def worker(paragraph_text):
    return process_paragraph(paragraph_text)

if __name__ == '__main__':
    # 加载单词列表和词义信息
    file_path = './CET6_2.json'  # 假设这是包含JSON数据的文件
    headwords, word_info_dict = extract_headwords_and_info_from_file(file_path)

    # 打印所有提取到的headWord
    for i, word in enumerate(headwords, 1):
        print(f"HeadWord {i}: {word}")

    # 单词列表
    word_list = headwords

    # 转换为小写，便于匹配
    word_list = [word.lower() for word in word_list]
    word_set = set(word_list)  # 在主进程中定义 word_set，便于后续计算

    # 加载文档
    document = Document('SorcererStone.docx')
    print("document loaded")
    # 提取所有段落的文本
    paragraphs = [para.text for para in document.paragraphs]

    # 创建多进程池
    pool = multiprocessing.Pool(initializer=init_pool, initargs=(word_list, word_info_dict))

    # 并行处理段落
    results = pool.map(worker, paragraphs)

    pool.close()
    pool.join()

    print("search finish")
    # 创建新的文档
    new_document = Document()

    # 汇总所有找到的单词
    total_found_words = set()

    # 重新组装处理后的段落
    for result in results:
        runs, found_words = result  # 解包返回值
        total_found_words.update(found_words)  # 更新总的找到的单词集合
        para = new_document.add_paragraph()
        for run_info in runs:
            run = para.add_run(run_info['text'])
            if run_info.get('highlight', False):
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW

    # 保存新的文档
    new_document.save('highlighted_document-totle-2.docx')

    # 计算未找到的单词
    found_words_list = list(total_found_words)
    not_found_words = list(word_set - total_found_words)
 
    # 输出找到和未找到的单词数量
    print(f"找到的单词数量：{len(found_words_list)}")
    print(f"未找到的单词数量：{len(not_found_words)}")

    # 保存结果到文件
    with open('found_words.txt', 'w', encoding='utf-8') as f:
        for word in sorted(found_words_list):
            f.write(word + '\n')

    with open('not_found_words.txt', 'w', encoding='utf-8') as f:
        for word in sorted(not_found_words):
            f.write(word + '\n')
